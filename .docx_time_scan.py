from docx import Document
from pathlib import Path
for f in Path('docs').glob('*.docx'):
    if f.name.startswith('~$'):
        continue
    d = Document(f)
    ps = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                ps.extend(cell.paragraphs)
    found = [p for p in ps if '08:00' in p.text or 'trước 8h' in p.text.lower()]
    if found:
        print(f'### {f.name}')
        for p in found:
            print(repr(p.text))
            print('RUNS:', [r.text for r in p.runs])
